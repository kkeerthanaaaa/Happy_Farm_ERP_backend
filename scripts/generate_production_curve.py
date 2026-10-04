#!/usr/bin/env python3
"""
generate_production_curve.py
============================
Generates the authoritative Production Curve Excel export (.xlsx) from JSON input.
TABLE-ONLY FORMAT: ZERO CHARTS.
Reproduces the exact table structure, rows, row colors, and week-start column highlighting
from the reference production workbook.

Called by: backend/src/services/reports.service.ts  generateProductionCurveExport()
Usage:     python generate_production_curve.py <input_json_path> <output_xlsx_path>

Rows:
  1. DATE
  2. WEEKS
  3. NO.OF BIRDS
  4. PRODUCTION
  5. SELECTION
  6. SELECTION %
  7. DAMAGE/REJECTED
  8. Mortality
  9. Temp
  10. Feed Kgs
  11. Feed Gms/Bird
  12. STD %
  13. ACT %
"""

import sys
import os
import json
import math
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ────────────────────────────────────────────────────────
# Reference Standard Production Curves (from reference workbook CF STD / FR STD)
# ────────────────────────────────────────────────────────

CF_STD_CURVE = {
    16: 0.0, 17: 0.0, 18: 0.0, 19: 3.19, 20: 17.96, 21: 46.5, 22: 73.82, 23: 85.25, 24: 87.46,
    25: 90.71, 26: 89.11, 27: 89.24, 28: 88.83, 29: 89.04, 30: 90.37, 31: 90.42, 32: 90.93,
    33: 90.27, 34: 88.83, 35: 91.19, 36: 89.13, 37: 89.78, 38: 90.26, 39: 89.7, 40: 87.02,
    41: 88.63, 42: 88.03, 43: 88.02, 44: 87.26, 45: 87.56, 46: 87.02, 47: 86.01, 48: 86.56,
    49: 86.51, 50: 85.95, 51: 85.54, 52: 85.03, 53: 84.56, 54: 84.0, 55: 84.23, 56: 84.23,
    57: 83.21, 58: 83.15, 59: 83.02, 60: 82.0, 61: 82.58, 62: 82.54, 63: 82.01, 64: 81.01,
    65: 81.25, 66: 81.75, 67: 81.54, 68: 81.58, 69: 80.47, 70: 80.45, 71: 80.98, 72: 80.56,
    73: 79.54, 74: 79.54, 75: 79.25, 76: 78.25, 77: 78.85, 78: 78.24, 79: 78.01, 80: 78.14
}

FR_STD_CURVE = {
    16: 0.0, 17: 0.0, 18: 0.0, 19: 0.0, 20: 0.0, 21: 3.54, 22: 18.66, 23: 46.59, 24: 70.4,
    25: 82.31, 26: 83.2, 27: 84.83, 28: 89.87, 29: 90.06, 30: 89.25, 31: 89.4, 32: 87.25,
    33: 87.97, 34: 86.65, 35: 86.0, 36: 86.21, 37: 85.01, 38: 85.21, 39: 84.01, 40: 84.21,
    41: 83.21, 42: 82.1, 43: 80.53, 44: 84.64, 45: 81.64, 46: 82.13, 47: 79.83, 48: 80.72,
    49: 82.6, 50: 83.49, 51: 83.5, 52: 83.57, 53: 84.59, 54: 85.59, 55: 85.87, 56: 86.25,
    57: 86.98, 58: 85.26, 59: 85.4, 60: 84.26, 61: 84.01, 62: 84.7, 63: 84.25, 64: 83.02,
    65: 83.41, 66: 83.21, 67: 83.04, 68: 82.01, 69: 82.01, 70: 82.65, 71: 82.98, 72: 82.1,
    73: 81.02, 74: 81.54, 75: 80.29, 76: 80.25, 77: 79.02, 78: 79.54, 79: 78.01, 80: 78.2
}

def get_std_pct(curve_type, age_weeks):
    """
    Interpolate standard production % using daily increments matching reference workbook.
    age_weeks is in W.D format (e.g. 21.0, 21.1, ..., 21.6).
    Daily increment = (Std(W+1) - Std(W)) / 7.
    Std(W.D) = Std(W) + D * daily_increment.
    """
    curve = CF_STD_CURVE if curve_type == 'CF_STD' else FR_STD_CURVE
    w_int = int(age_weeks)
    d_frac = round((age_weeks - w_int) * 10)  # 0 to 6
    if d_frac < 0 or d_frac > 6:
        d_frac = min(6, max(0, round((age_weeks - w_int) * 7)))

    weeks_list = sorted(curve.keys())
    if w_int <= weeks_list[0]:
        base_val = curve[weeks_list[0]]
        next_val = curve.get(weeks_list[0] + 1, base_val)
    elif w_int >= weeks_list[-1]:
        return curve[weeks_list[-1]]
    else:
        base_val = curve.get(w_int, 0.0)
        next_val = curve.get(w_int + 1, base_val)

    inc = (next_val - base_val) / 7.0
    val = base_val + d_frac * inc
    return round(max(0.0, val), 2)


# ────────────────────────────────────────────────────────
# Styles Matching Reference Production Sheet
# ────────────────────────────────────────────────────────

FILL_WEEK_HEADER = PatternFill(start_color='F8CBAD', end_color='F8CBAD', fill_type='solid')  # Solid Peach for week-start cell
FILL_WEEK_COL    = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')  # Soft Peach for week-start column
FILL_PURPLE      = PatternFill(start_color='7030A0', end_color='7030A0', fill_type='solid')  # Deep Purple for SELECTION %
FILL_CYAN        = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')  # Light Blue for Feed Gms/Bird
FILL_YELLOW      = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')  # Yellow for ACT %
FILL_WHITE       = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

FONT_BOLD_BLACK  = Font(name='Calibri', size=11, bold=True, color='000000')
FONT_REG_BLACK   = Font(name='Calibri', size=11, bold=False, color='000000')
FONT_BOLD_YELLOW = Font(name='Calibri', size=11, bold=True, color='FFFF00')
FONT_BOLD_RED    = Font(name='Calibri', size=11, bold=True, color='FF0000')

BORDER_THIN = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000'),
)
BORDER_WEEK_START = Border(
    left=Side(style='medium', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000'),
)

ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_LEFT   = Alignment(horizontal='left', vertical='center')

ROW_LABELS = [
    'DATE',
    'WEEKS',
    'NO.OF BIRDS',
    'PRODUCTION',
    'SELECTION',
    'SELECTION %',
    'DAMAGE/REJECTED',
    'Mortality',
    'Temp',
    'Feed Kgs',
    'Feed Gms/Bird',
    'STD %',
    'ACT %',
]


def build_farm_sheet(wb, farm):
    farm_name = farm.get('farmName', farm.get('farmId', 'Unknown'))
    reports = farm.get('reports', [])
    curve_type = farm.get('productionCurve', 'CF_STD')
    initial_birds = farm.get('initialBirds', 1200)

    # Sanitize sheet title (Excel max 31 chars)
    sheet_title = farm_name[:31]
    for ch in ['/', '\\', '*', '?', '[', ']', ':']:
        sheet_title = sheet_title.replace(ch, '-')
    if sheet_title in wb.sheetnames:
        sheet_title = sheet_title[:28] + '_' + str(len(wb.sheetnames))
    ws = wb.create_sheet(title=sheet_title)

    # ZERO CHARTS - purely the production table

    # Freeze panes: freeze Column A so metric labels stay visible on horizontal scroll
    ws.freeze_panes = 'B1'

    # Column A width (ample space to prevent clipping of 'DAMAGE/REJECTED' and 'Feed Gms/Bird')
    ws.column_dimensions['A'].width = 21.5

    # Populate Column A labels (Rows 1 to 13)
    for idx, label in enumerate(ROW_LABELS):
        r = idx + 1
        cell = ws.cell(row=r, column=1, value=label)
        cell.alignment = ALIGN_LEFT
        cell.border = BORDER_THIN
        ws.row_dimensions[r].height = 20.0

        if label == 'STD %':
            cell.font = FONT_BOLD_RED
            cell.fill = FILL_WHITE
        elif label == 'ACT %':
            cell.font = FONT_BOLD_BLACK
            cell.fill = FILL_YELLOW
        elif label == 'SELECTION %':
            cell.font = FONT_BOLD_BLACK
            cell.fill = FILL_WHITE
        elif label == 'Feed Gms/Bird':
            cell.font = FONT_BOLD_BLACK
            cell.fill = FILL_WHITE
        else:
            cell.font = FONT_BOLD_BLACK
            cell.fill = FILL_WHITE

    if not reports:
        return

    # Sort reports chronologically
    reports.sort(key=lambda r: r.get('submissionDate', ''))

    # Populate Data Columns (B onwards)
    for col_idx, rep in enumerate(reports):
        col = col_idx + 2
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 12.0

        date_str = rep.get('submissionDate', '')
        age_weeks = float(rep.get('ageWeeks', 0) or 0)
        birds = rep.get('openingBirdCount') or initial_birds
        prod = rep.get('eggsProduced')
        sel = rep.get('selectionEggs')
        mort = rep.get('mortality')
        temp = rep.get('temperature')
        feed = rep.get('feedKg')

        # Detect week start dynamically (whole number age: 20, 21, 22...)
        is_week_start = (round(age_weeks, 2) % 1 == 0) or (abs(age_weeks - round(age_weeks)) < 0.01)

        # Parse date to date object for Excel date format
        date_obj = None
        if date_str:
            try:
                date_obj = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
            except:
                date_obj = date_str

        # Calculations (same column aligned, zero-division safe)
        damage_rej = None
        if prod is not None and sel is not None:
            damage_rej = max(0, prod - sel)

        sel_pct = None
        if prod and prod > 0 and sel is not None:
            sel_pct = round((sel / prod) * 100, 2)

        feed_gms = None
        if birds and birds > 0 and feed is not None:
            feed_gms = int(round((feed * 1000) / birds))

        act_pct = None
        if birds and birds > 0 and prod is not None:
            act_pct = round((prod / birds) * 100, 2)

        std_pct = get_std_pct(curve_type, age_weeks)

        # Week display: integer for week start (21), decimal for day offset (21.1)
        week_val = int(round(age_weeks)) if is_week_start else round(age_weeks, 1)

        border = BORDER_WEEK_START if is_week_start else BORDER_THIN
        base_fill = FILL_WEEK_COL if is_week_start else FILL_WHITE

        # Values mapping for rows 1 to 13
        row_data = [
            # Row 1: DATE
            (date_obj, 'dd-mmm', FONT_BOLD_BLACK, base_fill),
            # Row 2: WEEKS
            (week_val, 'General', FONT_BOLD_BLACK, FILL_WEEK_HEADER if is_week_start else FILL_WHITE),
            # Row 3: NO.OF BIRDS
            (birds, '#,##0', FONT_REG_BLACK, base_fill),
            # Row 4: PRODUCTION
            (prod if prod is not None else '', '#,##0', FONT_REG_BLACK, base_fill),
            # Row 5: SELECTION
            (sel if sel is not None else '', '#,##0', FONT_REG_BLACK, base_fill),
            # Row 6: SELECTION %
            (sel_pct if sel_pct is not None else '', '0.00', FONT_BOLD_YELLOW, FILL_PURPLE),
            # Row 7: DAMAGE/REJECTED
            (damage_rej if damage_rej is not None else '', '#,##0', FONT_REG_BLACK, base_fill),
            # Row 8: Mortality (blank if 0 or None, matching reference sheet)
            (mort if mort and mort > 0 else '', 'General', FONT_REG_BLACK, base_fill),
            # Row 9: Temp
            (temp if temp is not None else '', '0', FONT_REG_BLACK, base_fill),
            # Row 10: Feed Kgs
            (feed if feed is not None else '', '#,##0', FONT_REG_BLACK, base_fill),
            # Row 11: Feed Gms/Bird
            (feed_gms if feed_gms is not None else '', '0', FONT_BOLD_BLACK, FILL_CYAN),
            # Row 12: STD %
            (std_pct, '0.00', FONT_BOLD_RED, base_fill),
            # Row 13: ACT %
            (act_pct if act_pct is not None else '', '0.00', FONT_BOLD_BLACK, FILL_YELLOW),
        ]

        for idx, (val, fmt, font, fill) in enumerate(row_data):
            r = idx + 1
            cell = ws.cell(row=r, column=col, value=val)
            cell.number_format = fmt
            cell.font = font
            cell.fill = fill
            cell.alignment = ALIGN_CENTER
            cell.border = border


def main():
    if len(sys.argv) < 3:
        print('Usage: python generate_production_curve.py <input.json> <output.xlsx>', file=sys.stderr)
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2]

    with open(input_path, 'r', encoding='utf-8') as f:
        payload = json.load(f)

    farms = payload.get('farms', [])
    if not farms:
        print('ERROR: No farms in input payload', file=sys.stderr)
        sys.exit(1)

    wb = Workbook()
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    for farm in farms:
        try:
            build_farm_sheet(wb, farm)
        except Exception as e:
            farm_name = farm.get('farmName', farm.get('farmId', 'Unknown'))
            print('WARNING: Failed to build sheet for {}: {}'.format(farm_name, e), file=sys.stderr)
            import traceback
            traceback.print_exc(file=sys.stderr)
            sheet_title = farm_name[:31]
            for ch in ['/', '\\', '*', '?', '[', ']', ':']:
                sheet_title = sheet_title.replace(ch, '-')
            if sheet_title not in wb.sheetnames:
                ws = wb.create_sheet(title=sheet_title)
                ws['A1'] = 'Error generating data for {}: {}'.format(farm_name, str(e))

    if len(wb.sheetnames) == 0:
        ws = wb.create_sheet(title='No Data')
        ws['A1'] = 'No farm data available for the selected date range.'

    wb.save(output_path)
    print('SUCCESS: Table-only workbook saved to {}'.format(output_path))


if __name__ == '__main__':
    main()
